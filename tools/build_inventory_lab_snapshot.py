from datetime import date, datetime
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = Path.home() / 'Downloads'
TODAY = date(2026, 7, 28)


def clean(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace('�', 'N/D').strip()


def status(value):
    if not value or value in ('N/D', 'NO VISIBLE'):
        return 'FALTA CADUCIDAD'
    try:
        if isinstance(value, datetime):
            value = value.date()
        elif not isinstance(value, date):
            value = str(value).split(' ')[0]
            parts = str(value).split('-')
            value = date(int(parts[0]), int(parts[1]), 1) if len(parts) == 2 else date.fromisoformat(str(value))
        return 'VENCIDO' if value < TODAY else 'VIGENTE'
    except (TypeError, ValueError):
        return 'REVISAR FECHA'


def row(tipo, nombre, marca='', fabricante='', referencia='', lote='', caducidad='', equipo='', fuente='', **extra):
    return {
        'Tipo': tipo, 'Nombre': nombre, 'Marca': marca, 'Fabricante': fabricante,
        'Referencia': referencia, 'Lote': lote, 'Caducidad': caducidad,
        'Equipo/área': equipo, 'Fuente': fuente, 'Estado': status(caducidad),
        'Cantidad': extra.get('cantidad', ''), 'Unidad/presentación': extra.get('unidad', ''),
        'Analito/prueba': extra.get('analito', ''), 'Perfil/paquete': extra.get('perfil', ''),
        'Control/calibrador': extra.get('control', ''),
        'Mínimo (-2SD)': extra.get('minimo', ''), 'Máximo (+2SD)': extra.get('maximo', ''),
        'Método INCCA': extra.get('metodo_incca', ''),
        'Volumen muestra': extra.get('volumen_muestra', ''),
        'Volumen R1': extra.get('volumen_r1', ''), 'Volumen R2': extra.get('volumen_r2', ''),
        'Datos faltantes': extra.get('faltantes', ''),
    }


def load_general_inventory():
    path = DOWNLOADS / 'Inventario de Reactivos de Laboratorio.xlsx'
    rows = []
    for values in list(load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:]:
        values = list(values) + [''] * 6
        if not any(values[:6]):
            continue
        marca, nombre, reference, lote, expiry, category = map(clean, values[:6])
        rows.append(row(
            'REACTIVO/KIT', nombre, marca, referencia=reference, lote=lote,
            caducidad=expiry, equipo=category, fuente=path.name,
            faltantes='Cantidad; unidad; precio; analito exacto; equipo; inserto; factura',
        ))
    unique = {}
    for item in rows:
        unique[(item['Marca'], item['Nombre'], item['Referencia'], item['Lote'])] = item
    return list(unique.values())


def fixed_inventory():
    items = load_general_inventory()
    for name, reference, lote, expiry, quantity in (
        ('Norma-iLyse 3', 'IC-21731', '2143070724', '2028-06-30', '1 L'),
        ('Norma-iSol 3', 'IC-41731', '4149050324', '2028-02-29', '1 L'),
        ('Norma-iDil 3', 'IC-11731', '1143090924', '2028-08-31', '10 L'),
    ):
        items.append(row(
            'REACTIVO', name, 'Norma Instruments', 'Norma Instruments', reference, lote,
            expiry, 'Icon 3 / Hematología', 'Captura física entregada 2026-07-28',
            cantidad=quantity, unidad='L', analito='Biometría hemática; diferencial 3 partes',
            faltantes='Cantidad operativa; precio; factura; inserto; fecha de apertura',
        ))
    for marca, name, reference, lote, expiry, presentation, area in (
        ('Mission / ACON', 'Urinalysis Reagent Strips 10U', 'U031-101', 'URS6030051', '2028-03-20', '10 parámetros', 'Uroanálisis'),
        ('FUJIFILM', 'FDC Auto Tips', '100493', '95213', '', '96 piezas x 6', 'Fuji Dri-Chem'),
        ('Patches', 'Cubreobjetos', '', '', '', '1000 piezas', 'Microscopia'),
        ('Velab', 'Portaobjetos', 'VE-P20', '', '', '50 piezas', 'Microscopia'),
        ('On-Call', 'Lancetas 30G', 'G124-10A', '2303012', '2028-03-05', 'Punción digital', 'Toma de muestra'),
        ('BMH', 'Tubo gel/activador tapa oro', 'BV13100Y', '25120102', '2028-12', '5 ml', 'Toma de muestra'),
        ('Genérica / Atyde México', 'Aguja multimuestra 22G x 1 1/2', '592122380', '20240209', '2029-02-18', 'Unidad', 'Toma de muestra'),
        ('BMH', 'Tubo K2 EDTA tapa lila', 'BV13750L', '25120103', '2028-12', '4 ml', 'Hematología'),
        ('Genérica', 'Gluco-X solución curva tolerancia', '', '25091241', '', '75 g / 250 ml', 'Curva de tolerancia'),
        ('Patches', 'Banditas adhesivas', '', '202504', '2030-04', 'Unidad', 'Toma de muestra'),
        ('Genérica', 'Aguja multimuestra azul 23G x 1 1/2', '', '20260202', '2031-02-01', 'Unidad', 'Toma de muestra'),
    ):
        items.append(row(
            'CONSUMIBLE', name, marca, referencia=reference, lote=lote, caducidad=expiry,
            equipo=area, fuente='Datos entregados por usuario 2026-07-28', unidad=presentation,
            faltantes='Cantidad física; precio; fecha compra; factura; unidad de control',
        ))
    return items


def technical_rows():
    controls = ['C1 (407501)', 'C2 (315001)', 'ACP3', 'ALBUMINA', 'ALK-POINT R', 'ALT-POINT R', 'AST-POINT R', 'BUN 2 R', 'CHOLE-POINT R', 'CREATININA WINER', 'CREAT-POINT', 'D BIL-POINT R', 'GGT-POINT R', 'GLUCO OXIDASA-POINT R', 'HDL COL-POINT', 'IRON-POINT R', 'LDH-POINT R', 'T PRO-POINT R', 'TBILI-POINT R', 'TRIG-POINT R', 'URICO-POINT R']
    calibrators = ['MULTICALIBRADOR (324803)', 'ALBUMINA', 'BUN 2 R', 'BUN POINT', 'CHOLE-POINT R', 'COLESTEROL POINT', 'CREATININA WINER', 'CREAT-POINT', 'D BIL-POINT R', 'GLUCO OXIDASA-POINT R', 'IRON-POINT R', 'T PRO-POINT R', 'TBILI-POINT R', 'TRIG-POINT R', 'URICO-POINT R', 'ACP3', 'GLUCOSA OXIDASA', 'LDH-POINT R', 'GLUCOSA', 'BUN', 'CREA', 'ACIDOUR', 'COLE', 'TRI']
    def technical(kind, name, source):
        return row(kind, name, 'Point Scientific', 'Point Scientific', equipo='INCCA / Química líquida', fuente=source, perfil='Química de 6 / Perfil hepático según aplique', faltantes='Confirmar alcance; código; presentación; lote; caducidad; cantidad; consumo por analito; inserto; factura')
    c1_limits = (
        ('ALBUMINA', '2.2', '2.8'),
        ('ALK-POINT R', '51.2', '95.2'),
        ('ALT-POINT R', '43.2', '65.2'),
        ('AST-POINT R', '49.6', '79.6'),
        ('BUN 2 R', '12', '16'),
        ('CHOLE-POIN R', '121.2', '149.2'),
        ('CREATININA WINER', '0.71', '1.31'),
        ('CREAT-POINT', '0.71', '1.31'),
        ('D BIL-POIN R', '0.3', '1.1'),
        ('GGT-POIN R', '35.8', '61.8'),
        ('GLUCO OXIDASA-POIN R', '82.6', '102.6'),
        ('HDL COL-POIN', '46', '86'),
        ('IRON-POIN R', '51.6', '79.6'),
        ('LDH-POIN R', '96.2', '148.2'),
        ('T PRO-POIN R', '3.5', '4.3'),
        ('TBILI-POIN R', '0.38', '1.18'),
        ('TRIG-POIN R', '70.2', '120.2'),
        ('URICO-POINT R', '4.08', '5.88'),
    )
    c2_limits = (
        ('ALBUMINA', '3.9', '4.8'),
        ('ALK-POINT R', '158', '296'),
        ('ALT-POINT R', '195', '254'),
        ('AST-POINT R', '198', '299'),
        ('BUN 2 R', '38', '47'),
        ('CHOLE-POIN R', '261', '324'),
        ('CREATININA WINER', '4.23', '5.77'),
        ('CREAT-POINT', '4.23', '5.77'),
        ('D BIL-POIN R', '4.2', '6.4'),
        ('GGT-POIN R', '106', '160'),
        ('GLUCO OXIDASA-POIN R', '249', '305'),
        ('HDL COL-POIN', '94', '182'),
        ('IRON-POIN R', '179', '271'),
        ('LDH-POIN R', '269', '414'),
        ('T PRO-POIN R', '5.8', '7.2'),
        ('TBILI-POIN R', '3.5', '5.7'),
        ('TRIG-POIN R', '155', '261'),
        ('URICO-POINT R', '7.8', '11'),
    )
    def control_rows(control, limits, source):
        return [row(
        'CONTROL_QC', f'Control {control} - {analito}', 'Point Scientific', 'Point Scientific',
        equipo='INCCA / Química líquida', fuente=source,
        analito=analito, control=control, minimo=minimo, maximo=maximo,
        perfil='Química de 6 / Perfil hepático según aplique',
        faltantes='Unidad del rango; lote; caducidad; cantidad física; frecuencia QC; inserto; factura',
        ) for analito, minimo, maximo in limits]
    c1_rows = control_rows('C1 (407501)', c1_limits, 'Control 1 INCCA entregado 2026-07-28')
    c2_rows = control_rows('C2 (315001)', c2_limits, 'Control 2 INCCA entregado 2026-07-28')
    method_path = DOWNLOADS / 'integra las ultimas 2 respuestas revisa que todo....xlsx'
    if not method_path.exists():
        raise FileNotFoundError(f'No se encontró la fuente de métodos INCCA: {method_path}')
    method_rows = []
    source_rows = list(load_workbook(method_path, read_only=True, data_only=True).active.iter_rows(values_only=True))[1:]
    for values in source_rows:
        values = list(values) + ['', '', '', '']
        method_name, sample_volume, r1_volume, r2_volume = map(clean, values[:4])
        if not method_name:
            continue
        match = re.search(r'\(([^()]*)\)\s*$', method_name)
        code = match.group(1).strip() if match else ''
        analyte = method_name[:match.start()].strip() if match else method_name
        method_rows.append(row(
            'REACTIVO/METODO', method_name, 'Point Scientific', 'Point Scientific',
            referencia=code, equipo='INCCA / Química líquida', fuente=method_path.name,
            analito=analyte, perfil='Química de 6 / Perfil hepático según aplique',
            metodo_incca=method_name, volumen_muestra=sample_volume,
            volumen_r1=r1_volume, volumen_r2=r2_volume,
            faltantes='Reactivo exacto; lote; caducidad; cantidad física; precio; factura; inserto',
        ))
    return (c1_rows + c2_rows + [technical('CONTROL_QC', x, 'CONTROLES.aof') for x in controls] +
            [technical('CALIBRADOR', x, 'CALIBRADORES.acf') for x in dict.fromkeys(calibrators)] +
            method_rows)


def write_workbook(path, items):
    wb = Workbook()
    wb.remove(wb.active)
    headers = list(items[0])
    groups = [('Inventario consolidado', items), ('Reactivos y kits', [x for x in items if x['Tipo'] in ('REACTIVO/KIT', 'REACTIVO')]), ('Consumibles', [x for x in items if x['Tipo'] == 'CONSUMIBLE']), ('Controles QC', [x for x in items if x['Tipo'] == 'CONTROL_QC']), ('Calibradores', [x for x in items if x['Tipo'] == 'CALIBRADOR']), ('Metodos INCCA', [x for x in items if x['Tipo'] == 'REACTIVO/METODO'])]
    for title, rows in groups:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for item in rows: ws.append([item[h] for h in headers])
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78'); cell.alignment = Alignment(wrap_text=True)
        for col in range(1, len(headers) + 1): ws.column_dimensions[get_column_letter(col)].width = min(max(len(headers[col-1]) + 2, 14), 35)
        if ws.max_row > 1:
            table = Table(displayName='T' + ''.join(c for c in title if c.isalnum())[:20], ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)
    summary = wb.create_sheet('Resumen', 0)
    summary.append(['Indicador', 'Resultado'])
    summary.append(['Fecha de corte', str(TODAY)])
    summary.append(['Registros consolidados', len(items)])
    for kind in ('REACTIVO/KIT','REACTIVO','CONSUMIBLE','CONTROL_QC','CALIBRADOR','REACTIVO/METODO'):
        summary.append([kind, sum(x['Tipo'] == kind for x in items)])
    summary.append(['Vencidos', sum(x['Estado'] == 'VENCIDO' for x in items)])
    summary.append(['Sin caducidad', sum(x['Estado'] == 'FALTA CADUCIDAD' for x in items)])
    summary.append(['Pendientes de integración/confirmación', sum(bool(x['Datos faltantes']) for x in items)])
    summary.append(['Regla', 'No importar como lote liberado hasta completar cantidad, unidad, lote, caducidad y validación'])
    summary.column_dimensions['A'].width=38; summary.column_dimensions['B'].width=100
    for cell in summary[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
    wb.save(path)


if __name__ == '__main__':
    output = ROOT / 'docs' / 'manual' / 'INVENTARIO_LABORATORIO_CONSOLIDADO_2026-07-28.xlsx'
    write_workbook(output, fixed_inventory() + technical_rows())
    print(output)
